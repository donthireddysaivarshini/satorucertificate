"""Reading participant sheets (.xlsx / .csv) and checking them."""
import csv
import io
import re
from datetime import date, datetime

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

COLUMN_ALIASES = {
    "name": ["name", "full_name", "student_name", "participant_name", "participant"],
    "email": ["email", "email_id", "e_mail", "mail", "email_address"],
    "cert_no": ["certificate_number", "certificate_no", "cert_no", "certificate", "cert", "certificate_id", "serial_no", "serial"],
}


def norm(key):
    return re.sub(r"[^a-z0-9]+", "_", str(key).strip().lower()).strip("_")


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return f"{v.day} {v.strftime('%B %Y')}"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(fileobj, filename):
    """Return (header keys, list of (row number, {key: value}))."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(fileobj, read_only=True, data_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        raw = fileobj.read()
        text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Upload an .xlsx or .csv file.")
    if not rows:
        return [], []
    keys = [norm(h) if h is not None else "" for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        rec = {k: _cell(v) for k, v in zip(keys, r) if k}
        if any(rec.values()):
            out.append((i, rec))
    return [k for k in keys if k], out


def cert_number(raw, index, fmt, start):
    """(serial int or None, full certificate number) for one row."""
    fmt = fmt or "{n}"
    raw = (raw or "").strip()
    if raw.isdigit():
        n = int(raw)
    elif raw:  # already a full number such as SF/STW/012/2026
        parts = re.split(r"\{n(?::[^}]*)?\}", fmt)
        m = re.fullmatch(r"(\d+)".join(re.escape(p) for p in parts), raw) if len(parts) == 2 else None
        return (int(m.group(1)) if m else None), raw
    elif start is not None:
        n = int(start) + index
    else:
        return None, ""
    try:
        return n, fmt.format(n=n)
    except (ValueError, KeyError, IndexError):
        return n, f"{fmt}-{n}"


def parse(fileobj, filename, fmt, start, used_placeholders=(), existing_cert_nos=frozenset()):
    """
    Parse a sheet into certificate rows plus problems.
    Returns (rows, errors, warnings, columns). Each row: {row, name, email, cert_no, serial, data}.
    """
    columns, raw_rows = read_rows(fileobj, filename)
    errors, warnings, rows = [], [], []
    if not raw_rows:
        return [], ["The sheet has no data rows."], [], columns

    def pick(rec, canon):
        for k in COLUMN_ALIASES[canon]:
            if rec.get(k):
                return rec[k]
        return ""

    if not any(k in columns for k in COLUMN_ALIASES["name"]):
        errors.append(f"No Name column found. Columns seen: {', '.join(columns) or 'none'}.")
    if not any(k in columns for k in COLUMN_ALIASES["email"]):
        errors.append("No Email column found.")
    has_cert_col = any(k in columns for k in COLUMN_ALIASES["cert_no"])
    if not has_cert_col and start is None:
        errors.append("The sheet has no Certificate No column: set 'Number from' to auto-number.")

    # A placeholder like {college} can be filled from a column such as "College Name".
    aliases = {}
    for ph in used_placeholders:
        if ph in columns:
            continue
        matches = [c for c in columns if c.startswith(ph + "_") or c.endswith("_" + ph)]
        if len(matches) == 1:
            aliases[ph] = matches[0]
    columns = columns + list(aliases)

    seen_cert, seen_email = {}, {}
    known = {"name", "email", "cert_no", "serial", "first_name", "code", "verify_url", "event", "from_name"}
    for idx, (row_no, rec) in enumerate(raw_rows):
        name = pick(rec, "name")
        if not name:
            warnings.append(f"Row {row_no}: no name, skipped.")
            continue
        email = pick(rec, "email")
        n, cert_no = cert_number(pick(rec, "cert_no"), len(rows), fmt, start)
        who = f"Row {row_no} ({name})"
        if not email:
            errors.append(f"{who}: no email address.")
        elif not EMAIL_RE.match(email):
            errors.append(f"{who}: invalid email '{email}'.")
        if not cert_no:
            errors.append(f"{who}: no certificate number.")
        elif cert_no in seen_cert:
            errors.append(f"{who}: certificate no {cert_no} is also on row {seen_cert[cert_no]}.")
        elif cert_no in existing_cert_nos:
            errors.append(f"{who}: certificate no {cert_no} was already issued in another batch.")
        seen_cert.setdefault(cert_no, row_no)
        if email:
            e = email.lower()
            if e in seen_email:
                warnings.append(f"{who}: same email as row {seen_email[e]}.")
            seen_email.setdefault(e, row_no)
        data = {k: v for k, v in rec.items()}
        for ph, col in aliases.items():
            data[ph] = rec.get(col, "")
        for ph in used_placeholders:
            if ph not in known and ph in columns and not data.get(ph):
                warnings.append(f"{who}: '{ph}' is empty.")
        rows.append({"row": row_no, "name": name, "email": email, "cert_no": cert_no,
                     "serial": f"{n:03d}" if n is not None else "", "data": data})
    return rows, errors, warnings, columns
